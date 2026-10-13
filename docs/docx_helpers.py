import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), b_val.get('val', 'single'))
            node.set(qn('w:sz'), str(b_val.get('sz', 4)))
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), b_val.get('color', 'CCCCCC'))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def add_callout_box(doc, text, title="NOTE / SPECIFICATION", border_color="0284C7", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    b_spec = {'val': 'single', 'sz': 12, 'color': border_color}
    set_cell_borders(cell, left=b_spec, top=None, bottom=None, right=None)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{title}: ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(9.5)
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_box(doc, code_text, caption=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    b_spec = {'val': 'single', 'sz': 6, 'color': 'CBD5E1'}
    set_cell_borders(cell, left=b_spec, top=b_spec, bottom=b_spec, right=b_spec)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(6)
        cprun = cp.add_run(caption)
        cprun.font.name = "Calibri"
        cprun.font.size = Pt(9)
        cprun.italic = True
        cprun.font.color.rgb = RGBColor(71, 85, 105)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_diagram_box(doc, diagram_title, diagram_content, note=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    b_spec = {'val': 'single', 'sz': 8, 'color': '0284C7'}
    set_cell_borders(cell, left=b_spec, top=b_spec, bottom=b_spec, right=b_spec)
    
    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    run_t = p0.add_run(diagram_title)
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(2, 132, 199)
    
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.05
    run_c = p1.add_run(diagram_content)
    run_c.font.name = "Consolas"
    run_c.font.size = Pt(8.5)
    run_c.font.color.rgb = RGBColor(15, 23, 42)
    
    if note:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(4)
        p2.paragraph_format.space_after = Pt(0)
        run_n = p2.add_run(note)
        run_n.italic = True
        run_n.font.name = "Calibri"
        run_n.font.size = Pt(9)
        run_n.font.color.rgb = RGBColor(71, 85, 105)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_result_placeholder(doc, fig_table_id, title, x_axis, y_axis, params, measurement_reqs):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FFFBEB")  # Amber 50
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    b_spec = {'val': 'single', 'sz': 10, 'color': 'D97706'}  # Amber 600
    set_cell_borders(cell, left=b_spec, top=b_spec, bottom=b_spec, right=b_spec)
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(3)
    r0 = p0.add_run(f"[RESULTS TO BE INSERTED AFTER EXPERIMENTS: {fig_table_id}]")
    r0.bold = True
    r0.font.name = "Calibri"
    r0.font.size = Pt(10)
    r0.font.color.rgb = RGBColor(180, 83, 9)
    
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(f"Title: {title}\n")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(15, 23, 42)
    
    details = (
        f"• Independent Variable (X-Axis): {x_axis}\n"
        f"• Dependent Variable (Y-Axis): {y_axis}\n"
        f"• Workload & Architectural Parameters: {params}\n"
        f"• Measurement & Execution Protocol: {measurement_reqs}\n"
        f"• Status: Pending physical hardware execution / cycle-accurate simulation run."
    )
    r2 = p1.add_run(details)
    r2.font.name = "Calibri"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Helper functions defined successfully.")
